"""Ingest every file in the `information/` folder into data.csv and a search index.

For each file the script extracts text based on its type:
  - PDFs                -> pdfplumber
  - Excel spreadsheets  -> openpyxl (.xlsx/.xlsm) or xlrd (.xls, optional)
  - Word documents      -> python-docx
  - Plain text / minutes-> read directly
  - Images              -> OCR via EasyOCR (pure pip install, no native binary)

Each file becomes one row in data.csv (summary + keywords) AND its full text is
stored in search_index.db so every word in every document is searchable.

Keywords use an LLM when GROQ_API_KEY or GOOGLE_API_KEY/GEMINI_API_KEY is set
(see llm.py) and fall back to a local extractor otherwise - the pipeline never
hard-fails on a missing key or network. NOTE: with a key set, the first ~6000
characters of each document are sent to that provider; use --no-ai to avoid it.

Usage:
  python ingest.py              # incremental, AI keywords if a key is set
  python ingest.py --rebuild    # ignore the cache, reprocess everything
  python ingest.py --no-ai      # local keyword extraction only (nothing leaves your PC)
"""
from __future__ import annotations

import argparse
import collections
import csv
import datetime
import os
import re
import tempfile
from contextlib import suppress
from pathlib import Path

import llm
import searchlib
from config import (CSV_FILE, DB_FILE, EXCEL_EXTS, FIELDNAMES, IMAGE_EXTS, INFORMATION_DIR,
                    LLM_INPUT_CHARS, MAX_KEYWORDS, MAX_PDF_PAGES, MAX_SUMMARY, MAX_TEXT_CHARS,
                    MAX_TEXT_FILE_BYTES, TEXT_EXTS)

CACHEABLE_STATUS = {"ok", "no text"}   # statuses that don't need to be retried every run
_EASYOCR_READER = None
_EASYOCR_ERROR = None


# --- Classification & text reduction --------------------------------------
def classify(ext: str) -> str:
    """Map a file extension to a human-friendly document category."""
    ext = ext.lower()
    if ext == ".pdf":
        return "PDF"
    if ext in EXCEL_EXTS:
        return "Spreadsheet"
    if ext == ".docx":
        return "Document"
    if ext in IMAGE_EXTS:
        return "Image"
    if ext in TEXT_EXTS:
        return "Text"
    return "Other"


def infer_category(filename: str, text: str) -> str:
    combined = (filename + " " + text[:200]).lower()
    if re.search(r"\bsop\b|standard operating proc", combined):
        return "SOP"
    if re.search(r"\bcircular\b|\bsurat pekeliling\b|\bsurat edaran\b", combined):
        return "Circular"
    if re.search(r"\bpolicy\b|\bpolicies\b|\bdasar\b", combined):
        return "Policy"
    if re.search(r"\bguideline\b|\bguidelines\b|\bpanduan\b|\bgaris panduan\b", combined):
        return "Guideline"
    if re.search(r"\breport\b|\breports\b|\blaporan\b", combined):
        return "Report"
    if re.search(r"\bminutes\b|\bmeeting minutes\b|\bminit mesyuarat\b|\bminit\b", combined):
        return "Minutes"
    return "Other"


def summarize_text(text: str) -> str:
    """Reduce extracted text to a single representative sentence."""
    if not text:
        return ""
    text = " ".join(text.split())
    summary = text
    for sentence in re.split(r"(?<=[.!?])\s+", text):
        candidate = sentence.strip()
        if len(candidate) >= 15 and re.search(r"[A-Za-z]", candidate):
            summary = candidate
            break
    if len(summary) > MAX_SUMMARY:
        summary = summary[: MAX_SUMMARY - 3].rstrip() + "..."
    return summary


# --- Keyword extraction (AI with local fallback) ---------------------------
_STOPWORDS = frozenset("""
a an the and or but if then else for to of in on at by with from as is are was were
be been being this that these those it its it's into over under about above below
not no nor so than too very can will just don should now i you he she we they them
me my your our their his her out up down off again further once here there all any
both each few more most other some such only own same which who whom what
when where why how do does did doing have has had having would could may might must
page version date name number total based using use used shall per via etc
""".split())


def _parse_keywords(raw: str) -> list[str]:
    """Normalize an LLM/text response into a clean keyword list."""
    seen, out = set(), []
    for part in re.split(r"[,\n;]+", raw):
        kw = re.sub(r"^[\s\-\*\d\.\)]+", "", part).strip().strip("\"'").lower()
        if 2 <= len(kw) <= 40 and kw not in seen:
            seen.add(kw)
            out.append(kw)
        if len(out) >= MAX_KEYWORDS:
            break
    return out


def _keywords_via_llm(text: str) -> list[str] | None:
    prompt = ("Extract the most important keywords and key phrases from the document text below. "
              f"Return at most {MAX_KEYWORDS} items as a comma-separated list, lowercase, "
              "no numbering, no extra commentary.\n\nTEXT:\n" + text[:LLM_INPUT_CHARS])
    raw = llm.complete(prompt, max_tokens=120)
    return _parse_keywords(raw) if raw else None


def _keywords_local(text: str) -> list[str]:
    """Dependency-free fallback: frequency-ranked content words and bigrams."""
    tokens = [t for t in re.findall(r"[A-Za-z][A-Za-z\-']+", text.lower())
              if len(t) >= 3 and t not in _STOPWORDS]
    if not tokens:
        return []
    freq = collections.Counter(tokens)
    for a, b in zip(tokens, tokens[1:]):
        freq[f"{a} {b}"] += 2                       # weight phrases higher
    out: list[str] = []
    for term, _ in sorted(freq.items(), key=lambda kv: (-kv[1], kv[0])):
        if any(term != kw and term in kw.split() for kw in out):
            continue                                # skip a word already inside a chosen phrase
        out.append(term)
        if len(out) >= MAX_KEYWORDS:
            break
    return out


def extract_keywords(text: str, use_ai: bool) -> tuple[str, str]:
    """Return (keywords_csv, method) where method is none / ai / local / local (ai failed)."""
    if not text or not text.strip():
        return "", "none"
    if use_ai and llm.get_client():
        kws = _keywords_via_llm(text)
        if kws:
            return ", ".join(kws), "ai"
        return ", ".join(_keywords_local(text)), "local (ai failed)"
    return ", ".join(_keywords_local(text)), "local"


# --- Extractors: each returns (extracted_text, status_message) ---------------
def extract_pdf(path: str):
    try:
        import pdfplumber
    except ImportError:
        return "", "skipped: pdfplumber not installed (pip install pdfplumber)"
    try:
        parts = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages[:MAX_PDF_PAGES]:
                try:   # removes the "RReeggiissttrraattiioonn" doubling from fake-bold PDFs
                    text = page.dedupe_chars(tolerance=1).extract_text()
                except Exception:  # noqa: BLE001 - older pdfplumber or odd page
                    text = page.extract_text()
                parts.append(text or "")
        text = "\n".join(parts)
        return text, ("ok" if text.strip() else "no text")   # "no text" usually means a scanned PDF
    except Exception as e:  # noqa: BLE001
        return "", f"error: {e}"


def extract_excel(path: str):
    if path.lower().endswith(".xls"):
        try:
            import xlrd
        except ImportError:
            return "", "skipped: .xls needs xlrd (pip install xlrd)"
        try:
            book, parts = xlrd.open_workbook(path), []
            for sheet in book.sheets():
                parts.append(f"[Sheet: {sheet.name}]")
                for r in range(sheet.nrows):
                    cells = [str(c) for c in sheet.row_values(r) if str(c).strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts), "ok"
        except Exception as e:  # noqa: BLE001
            return "", f"error: {e}"
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "", "skipped: openpyxl not installed (pip install openpyxl)"
    wb = None
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts), "ok"
    except Exception as e:  # noqa: BLE001
        return "", f"error: {e}"
    finally:
        if wb is not None:
            wb.close()


def extract_docx(path: str):
    try:
        import docx
    except ImportError:
        return "", "skipped: python-docx not installed (pip install python-docx)"
    try:
        doc = docx.Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts), "ok"
    except Exception as e:  # noqa: BLE001
        return "", f"error: {e}"


def extract_text(path: str):
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return f.read(MAX_TEXT_FILE_BYTES), "ok"
    except Exception as e:  # noqa: BLE001
        return "", f"error: {e}"


def _get_easyocr_reader():
    """Lazily build and cache a single EasyOCR reader (English)."""
    global _EASYOCR_READER, _EASYOCR_ERROR
    if _EASYOCR_READER is not None or _EASYOCR_ERROR is not None:
        return _EASYOCR_READER
    try:
        import easyocr
    except ImportError:
        _EASYOCR_ERROR = "skipped: easyocr not installed (pip install easyocr)"
        return None
    try:
        _EASYOCR_READER = easyocr.Reader(["en"], gpu=False, verbose=False)
    except Exception as e:  # noqa: BLE001
        _EASYOCR_ERROR = f"error building EasyOCR reader: {e}"
    return _EASYOCR_READER


def extract_image(path: str):
    reader = _get_easyocr_reader()
    if reader is None:
        return "", _EASYOCR_ERROR or "skipped: OCR unavailable"
    try:
        text = " ".join(reader.readtext(path, detail=0, paragraph=True))
        return text, ("ok" if text.strip() else "no text")
    except Exception as e:  # noqa: BLE001
        return "", f"error: {e}"


def extract(path: str, ext: str):
    # NOTE: .doc (old Word 97-2003 format) is NOT supported. It requires a native
    # binary (antiword or LibreOffice). Convert .doc files to .docx before ingesting.
    ext = ext.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext in EXCEL_EXTS:
        return extract_excel(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext in IMAGE_EXTS:
        return extract_image(path)
    if ext in TEXT_EXTS:
        return extract_text(path)
    return "", "skipped: unsupported file type"


# --- Incremental cache ------------------------------------------------------
def _stamp(stat) -> str:
    return datetime.datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S")


def _load_cache(out) -> dict:
    """Index existing CSV rows by relative_path for incremental reuse."""
    if not os.path.exists(out):
        return {}
    try:
        with open(out, newline="", encoding="utf-8") as f:
            return {r["relative_path"]: r for r in csv.DictReader(f) if r.get("relative_path")}
    except Exception:  # noqa: BLE001 - a bad/old cache is simply ignored
        return {}


def _cache_hit(cached: dict | None, stat) -> bool:
    """A cached row is reusable if size + mtime match and extraction had succeeded."""
    return bool(cached) and cached.get("status") in CACHEABLE_STATUS \
        and str(stat.st_size) == str(cached.get("size_bytes")) and cached.get("modified") == _stamp(stat)


# --- Pipeline --------------------------------------------------------------
def iter_files(root):
    for dirpath, dirs, files in os.walk(root):
        dirs[:] = sorted(d for d in dirs if not d.startswith("."))
        for name in sorted(files):
            if name.lower() == "readme.txt" or name.startswith((".", "~$")):
                continue
            yield os.path.join(dirpath, name)


def _atomic_write(out, rows: list[dict]) -> None:
    """Write CSV via a temp file then replace, so a locked target fails cleanly."""
    out = str(out)
    fd, tmp = tempfile.mkstemp(prefix=".ingest-", suffix=".csv", dir=os.path.dirname(os.path.abspath(out)) or ".")
    try:
        with os.fdopen(fd, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
            writer.writeheader()
            writer.writerows(rows)
        os.replace(tmp, out)                        # atomic on the same volume
    except PermissionError as e:
        with suppress(FileNotFoundError):
            os.remove(tmp)
        raise PermissionError(f"Cannot write '{out}' - it may be open in another program "
                              f"(e.g. Excel). Close it and re-run. ({e})") from e
    except Exception:
        with suppress(FileNotFoundError):
            os.remove(tmp)
        raise


def ingest(root=INFORMATION_DIR, out=CSV_FILE, db=DB_FILE, use_ai=True, rebuild=False, log=print):
    """Scan `root`, update the CSV catalog and the full-text index. Returns (rows, stats)."""
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    cache = {} if rebuild else _load_cache(out)
    files = list(iter_files(root))
    rows, stats = [], collections.Counter()
    conn = searchlib.connect(db)
    try:
        for doc_id, path in enumerate(files, start=1):
            ext = os.path.splitext(path)[1]
            rel = os.path.relpath(path, root).replace("\\", "/")
            stat = os.stat(path)
            stamp = _stamp(stat)
            cached = cache.get(rel)

            if _cache_hit(cached, stat) and searchlib.is_current(conn, rel, stat.st_size, stamp):
                row = dict(cached, doc_id=doc_id)
                if not row.get("category"):
                    row["category"] = infer_category(
                        os.path.basename(path),
                        cached.get("summary", "") + " " + cached.get("keywords", "")
                    )
                rows.append(row)
                stats["reused"] += 1
                continue

            log(f"[{doc_id}/{len(files)}] {rel}")
            text, status = extract(path, ext)
            text = text[:MAX_TEXT_CHARS]
            summary = summarize_text(text)
            if _cache_hit(cached, stat) and cached.get("keywords") and cached.get("summary") == summary:
                keywords, method = cached["keywords"], "cached"      # only the index was missing
            else:
                keywords, method = extract_keywords(text, use_ai=use_ai)
            stats["keywords: " + method] += 1
            stats["extracted ok" if status in CACHEABLE_STATUS else "problems"] += 1

            category = infer_category(os.path.basename(path), text)
            rows.append({"doc_id": doc_id, "file_name": os.path.basename(path), "relative_path": rel,
                         "doc_type": classify(ext), "size_bytes": stat.st_size, "modified": stamp,
                         "status": status, "summary": summary, "keywords": keywords, "category": category})
            title = f"{Path(path).stem.replace('_', ' ')} {keywords} {category}"
            searchlib.store(conn, rel, classify(ext), title, text, stat.st_size, stamp)

        stats["removed from index"] = searchlib.prune(conn, {r["relative_path"] for r in rows})
    finally:
        conn.close()
    _atomic_write(out, rows)
    return rows, stats


def main(argv=None):
    ap = argparse.ArgumentParser(description="Build data.csv and the full-text search index.")
    ap.add_argument("--rebuild", action="store_true", help="ignore the cache and reprocess every file")
    ap.add_argument("--no-ai", action="store_true", help="local keyword extraction only (no data sent anywhere)")
    args = ap.parse_args(argv)

    rows, stats = ingest(use_ai=not args.no_ai, rebuild=args.rebuild)
    print(f"\nIngested {len(rows)} file(s) into {CSV_FILE.name} and {DB_FILE.name}")
    reused = stats.pop("reused", 0)
    print(f"  reused from cache: {reused}")
    for k, v in sorted(stats.items()):
        print(f"  {k}: {v}")
    if llm.last_error:
        print(f"  AI warning: {llm.last_error}")
    for r in rows:
        if r["status"] not in CACHEABLE_STATUS:
            print(f"    - {r['file_name']}: {r['status']}")
        elif r["status"] == "no text":
            print(f"    - {r['file_name']}: no extractable text (scanned? try OCR/images)")


if __name__ == "__main__":
    main()
